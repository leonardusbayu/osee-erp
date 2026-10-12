"""Private staged registration intake. No network fetches or financial posting."""

import csv
import hashlib
import io
import zipfile
from datetime import date, datetime

from django.core.exceptions import ValidationError
from django.db import transaction

from finance import services as finance
from finance.models import Product

from .models import IntakeBatch, IntakeRow, Order, Registration
from .services import (
    _event,
    _org,
    _save,
    _text,
    _version,
    add_order_line,
    register_participant,
)

FIELDS = {
    "external_id": "ID respons / pendaftaran",
    "name": "Nama peserta",
    "email": "Email",
    "phone": "Nomor kontak",
    "planned_date": "Tanggal tes",
}
MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000


def _table(content, filename):
    if not isinstance(content, bytes) or not 1 <= len(content) <= MAX_BYTES:
        raise ValidationError("Pilih file CSV/XLSX berukuran maksimal 5 MB.")
    if filename.lower().endswith(".csv"):
        try:
            text = content.decode("utf-8-sig")
            dialect = csv.Sniffer().sniff(text[:8000], delimiters=",;\t")
        except UnicodeDecodeError as exc:
            raise ValidationError("Gunakan CSV UTF-8 atau XLSX.") from exc
        except csv.Error:
            dialect = csv.excel
        try:
            raw = []
            for row in csv.reader(io.StringIO(text), dialect):
                raw.append(row)
                if len(raw) > MAX_ROWS + 1:
                    raise ValidationError("Maksimal 5.000 baris per unggahan.")
        except csv.Error as exc:
            raise ValidationError("Struktur CSV tidak dapat dibaca.") from exc
    elif filename.lower().endswith(".xlsx"):
        from openpyxl import load_workbook

        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if (
                    len(archive.infolist()) > 300
                    or sum(x.file_size for x in archive.infolist()) > 30 * 1024 * 1024
                ):
                    raise ValidationError(
                        "Isi Excel melebihi batas aman. Pecah sumber menjadi beberapa file."
                    )
            workbook = load_workbook(
                io.BytesIO(content), read_only=True, data_only=False, keep_links=False
            )
            try:
                sheet = workbook.active
                if sheet.max_row > MAX_ROWS + 1 or sheet.max_column > 50:
                    raise ValidationError(
                        "Maksimal 5.000 baris dan 50 kolom per unggahan."
                    )
                raw = []
                for row in sheet.iter_rows():
                    if any(cell.data_type == "f" for cell in row):
                        raise ValidationError("Gunakan salinan nilai tanpa formula.")
                    raw.append([cell.value for cell in row])
            finally:
                workbook.close()
        except ValidationError:
            raise
        except Exception as exc:
            raise ValidationError("File XLSX tidak dapat dibaca.") from exc
    else:
        raise ValidationError("Gunakan file .csv atau .xlsx.")
    if len(raw) < 2 or any(len(row) > 50 for row in raw):
        raise ValidationError(
            "File memerlukan judul kolom pada baris pertama dan baris pendaftaran."
        )

    def cell(value):
        value = (
            value.isoformat()
            if isinstance(value, (date, datetime))
            else ""
            if value is None
            else str(value)
        )
        if len(value) > 2000 or "\x00" in value:
            raise ValidationError(
                "Satu sel terlalu panjang atau berisi karakter yang tidak valid."
            )
        return value

    raw = [[cell(value) for value in row] for row in raw]
    return [x[:120] or f"Kolom {i + 1}" for i, x in enumerate(raw[0])], raw[1:]


@transaction.atomic
def stage_upload(*, organization, actor, content, filename, source):
    org = _org(organization, actor)
    source, filename = _text(source, "Nama sumber"), _text(filename, "Nama file", 200)
    headers, rows = _table(content, filename)
    digest = hashlib.sha256(content).hexdigest()
    old = IntakeBatch.objects.filter(
        organization=org, source=source, digest=digest
    ).first()
    if old:
        return old
    batch = _save(
        IntakeBatch(
            organization=org,
            created_by=actor,
            source=source,
            filename=filename,
            digest=digest,
            headers=headers,
        )
    )
    for index, row in enumerate(rows, 2):
        _save(
            IntakeRow(
                organization=org,
                created_by=actor,
                batch=batch,
                row_number=index,
                raw=row,
            )
        )
    _event(
        org,
        actor,
        batch,
        "intake.staged",
        "Sumber privat disimpan untuk pemeriksaan",
        after={"rows": len(rows), "sha256": digest},
    )
    return batch


@transaction.atomic
def map_batch(*, organization, actor, batch, mapping, expected_version):
    org = _org(organization, actor)
    batch = finance._scoped(IntakeBatch, batch, org, lock=True)
    _version(batch, expected_version)
    if batch.rows.exclude(status="pending").exists():
        raise ValidationError(
            "Pemetaan tidak dapat diganti setelah sebagian baris diputuskan. Unggah revisi sebagai sumber baru."
        )
    if (
        not isinstance(mapping, dict)
        or set(mapping) - set(FIELDS)
        or not mapping.get("name")
    ):
        # String '0' is intentionally a valid first-column selection.
        raise ValidationError("Pilih kolom nama peserta dan periksa pemetaan.")
    chosen = [str(value) for value in mapping.values() if value != ""]
    if len(set(chosen)) != len(chosen) or any(
        not x.isdigit() or int(x) >= len(batch.headers) for x in chosen
    ):
        raise ValidationError(
            "Satu kolom hanya dapat dipetakan satu kali dan harus berada dalam file."
        )
    mapping = {key: str(value) for key, value in mapping.items() if value != ""}
    seen = set()
    for row in batch.rows.all():
        values = {
            key: row.raw[int(index)].strip() if int(index) < len(row.raw) else ""
            for key, index in mapping.items()
        }
        issues = []
        if not values.get("name"):
            issues.append("Nama peserta belum tersedia")
        if (
            len(values.get("name", "")) > 200
            or len(values.get("external_id", "")) > 160
        ):
            issues.append("Nama atau ID sumber terlalu panjang")
        if not values.get("external_id"):
            values["external_id"] = f"snapshot:{batch.digest}:row:{row.row_number}"
            values["snapshot_identity"] = True
        if values["external_id"] in seen:
            issues.append("ID sumber berulang dalam file; periksa duplikat")
        seen.add(values["external_id"])
        if values.get("planned_date"):
            original = values["planned_date"]
            parsed = None
            for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%dT%H:%M:%S"):
                try:
                    parsed = datetime.strptime(original, pattern).date()
                    break
                except ValueError:
                    continue
            if parsed is None:
                issues.append(
                    "Tanggal tes belum jelas; perbaiki pada sumber dan unggah revisi"
                )
            else:
                values["planned_date"] = str(parsed)
        row.normalized, row.issues = values, issues
        _save(row)
    batch.mapping, batch.version = mapping, batch.version + 1
    _save(batch)
    _event(
        org,
        actor,
        batch,
        "intake.mapped",
        "Kolom sumber dipetakan",
        after={"mapping": mapping},
    )
    return batch


@transaction.atomic
def accept_rows(
    *,
    organization,
    actor,
    batch,
    row_ids,
    order,
    product,
    expected_version,
    reason,
    confirm_identity=False,
    reviewed_duplicates=False,
):
    org = _org(organization, actor)
    batch = finance._scoped(IntakeBatch, batch, org, lock=True)
    _version(batch, expected_version)
    order, product = (
        finance._scoped(Order, order, org),
        finance._scoped(Product, product, org),
    )
    reason = _text(reason, "Catatan pemeriksaan", 2000)
    if (
        confirm_identity is not True
        or type(reviewed_duplicates) is not bool
        or not isinstance(row_ids, list)
        or not row_ids
        or len(row_ids) > MAX_ROWS
    ):
        raise ValidationError(
            "Pilih baris dan konfirmasikan identitas pemesan serta ID sumber."
        )
    if any(type(pk) is not int for pk in row_ids) or len(set(row_ids)) != len(row_ids):
        raise ValidationError("Pilihan baris tidak valid.")
    rows = list(batch.rows.filter(pk__in=row_ids, status="pending"))
    if len(rows) != len(row_ids) or any(
        row.issues or not row.normalized.get("name") for row in rows
    ):
        raise ValidationError(
            "Ada baris belum dipetakan, bermasalah, telah diputuskan, atau dari sumber lain."
        )
    existing = []
    new_rows = []
    selected_dates_by_name = {}
    for row in rows:
        data = row.normalized
        old = (
            Registration.objects.filter(
                organization=org, source=batch.source, external_id=data["external_id"]
            )
            .select_related("line", "participant")
            .first()
        )
        if old:
            if (
                old.line.order_id,
                old.line.product_id,
                old.participant.name,
                old.participant.email,
                old.participant.phone,
                str(old.planned_date or ""),
            ) != (
                order.pk,
                product.pk,
                data["name"],
                data.get("email", ""),
                data.get("phone", ""),
                data.get("planned_date", ""),
            ):
                raise ValidationError(
                    "ID sumber sudah terdaftar dengan data berbeda. Gunakan pemeriksaan koreksi sumber."
                )
            existing.append((row, old))
        else:
            # These rows have not reached Registration yet. Compare the new
            # selection as well as persisted registrations before creating any.
            name_key = data["name"].casefold()
            planned_date = data.get("planned_date", "")
            selected_dates = selected_dates_by_name.setdefault(name_key, set())
            same_selection = bool(selected_dates) and (
                not planned_date
                or "" in selected_dates
                or planned_date in selected_dates
            )
            possible = Registration.objects.filter(
                organization=org, participant__name__iexact=data["name"]
            )
            if planned_date:
                possible = possible.filter(planned_date=planned_date)
            if (same_selection or possible.exists()) and not reviewed_duplicates:
                raise ValidationError(
                    "Ada kandidat nama/tanggal yang sama pada pilihan baris atau pendaftaran tersimpan. Periksa duplikat lalu konfirmasikan bila ini pendaftaran berbeda; tanggal kosong belum dapat membedakan peserta."
                )
            selected_dates.add(planned_date)
            new_rows.append(row)
    if new_rows:
        # One source selection gets a distinct order component, keeping later
        # additions from silently changing an already-issued invoice.
        reference = f"intake-{batch.pk}-{batch.version}"
        line = add_order_line(
            organization=org,
            actor=actor,
            order=order,
            product=product,
            quantity=len(new_rows),
            reference=reference,
        )
        for row in new_rows:
            data = row.normalized
            registration = register_participant(
                organization=org,
                actor=actor,
                line=line,
                name=data["name"],
                source=batch.source,
                external_id=data["external_id"],
                email=data.get("email", ""),
                phone=data.get("phone", ""),
                planned_date=data.get("planned_date"),
            )
            existing.append((row, registration))
    for row, registration in existing:
        row.registration, row.status = registration, "accepted"
        _save(row)
    batch.version += 1
    _save(batch)
    _event(
        org,
        actor,
        batch,
        "intake.accepted",
        reason,
        after={
            "row_ids": row_ids,
            "order_id": order.pk,
            "new": len(new_rows),
            "existing": len(existing) - len(new_rows),
            "duplicates_reviewed": reviewed_duplicates,
        },
    )
    return {"accepted": len(existing), "created": len(new_rows), "batch": batch}


@transaction.atomic
def reject_rows(*, organization, actor, batch, row_ids, expected_version, reason):
    org = _org(organization, actor)
    batch = finance._scoped(IntakeBatch, batch, org, lock=True)
    _version(batch, expected_version)
    reason = _text(reason, "Alasan tidak menggunakan baris", 2000)
    if (
        not isinstance(row_ids, list)
        or not row_ids
        or any(type(x) is not int for x in row_ids)
    ):
        raise ValidationError("Pilih baris yang akan diperiksa.")
    rows = list(batch.rows.filter(pk__in=set(row_ids), status="pending"))
    if len(rows) != len(set(row_ids)):
        raise ValidationError("Baris sudah diputuskan atau tidak tersedia.")
    for row in rows:
        row.status = "rejected"
        _save(row)
    batch.version += 1
    _save(batch)
    _event(org, actor, batch, "intake.rejected", reason, after={"row_ids": row_ids})
    return batch
