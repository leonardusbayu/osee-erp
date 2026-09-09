"""Bounded Excel/CSV parsing and explicit, repeatable BNI statement mapping."""
import csv
import hashlib
import io
import re
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.core.exceptions import ValidationError
from django.utils import timezone

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000
MAX_COLUMNS = 50
ALIASES = {
    "date": {"date", "tanggal", "tanggaltransaksi", "transactiondate", "tgltransaksi", "tgl"},
    "reference": {"reference", "referensi", "referensitransaksi", "noreferensi", "nomorreferensi", "transactionid", "transactionreference", "ref", "refno", "jurnal", "nojurnal"},
    "description": {"description", "keterangan", "uraian", "deskripsi", "transactiondescription", "details"},
    "amount": {"amount", "nominal", "jumlah", "nilai"},
    "debit": {"debit", "debet", "withdrawal", "debitamount", "mutasidebet", "mutasidebit"},
    "credit": {"credit", "kredit", "deposit", "creditamount", "mutasikredit"},
    "balance": {"balance", "saldo", "saldoberjalan", "runningbalance"},
    "currency": {"currency", "matauang", "ccy"},
}


def _header(value):
    return re.sub(r"[^a-z0-9]", "", str(value).strip().lower())


def suggest_mapping(headers):
    result = {}
    for field, names in ALIASES.items():
        found = next((str(i) for i, name in enumerate(headers) if _header(name) in names), "")
        result[field] = found
    return result


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, (int, float, Decimal)):
        return {"number": str(value)}
    text = str(value)
    if len(text) > 2000:
        raise ValidationError("Satu sel melebihi 2.000 karakter. Periksa file sumber.")
    return text


def read_statement(content, filename):
    if len(content) > MAX_BYTES or not content:
        raise ValidationError("Pilih file berukuran 1 byte sampai 5 MB.")
    extension = filename.rsplit(".", 1)[-1].lower()
    raw = []
    if extension == "csv":
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("cp1252")
        try:
            dialect = csv.Sniffer().sniff(text[:8000], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(io.StringIO(text), dialect)
        try:
            for row in reader:
                raw.append(row)
                if len(raw) > MAX_ROWS + 40:
                    raise ValidationError("Maksimum 5.000 transaksi per impor. Pecah file menurut periode.")
        except csv.Error as exc:
            raise ValidationError("Struktur CSV tidak dapat dibaca atau satu sel terlalu panjang.") from exc
    elif extension == "xlsx":
        from openpyxl import load_workbook
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                infos = archive.infolist()
                if len(infos) > 300 or sum(item.file_size for item in infos) > 30 * 1024 * 1024:
                    raise ValidationError("Ukuran isi Excel terlalu besar setelah dibuka.")
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
            try:
                sheet = workbook.active
                if sheet.max_row > MAX_ROWS + 40 or sheet.max_column > MAX_COLUMNS:
                    raise ValidationError("Maksimum 5.000 transaksi dan 50 kolom per impor.")
                for row in sheet.iter_rows():
                    if any(cell.data_type == "f" for cell in row):
                        raise ValidationError("File berisi formula. Gunakan unduhan bank asli atau salinan nilai tanpa formula.")
                    raw.append([cell.value for cell in row])
            finally:
                workbook.close()
        except ValidationError:
            raise
        except Exception as exc:
            raise ValidationError("Excel tidak dapat dibaca. Gunakan file .xlsx asli yang tidak terenkripsi.") from exc
    elif extension == "xls":
        import xlrd
        try:
            workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
            try:
                sheet = workbook.sheet_by_index(0)
                if sheet.nrows > MAX_ROWS + 40 or sheet.ncols > MAX_COLUMNS:
                    raise ValidationError("Maksimum 5.000 transaksi dan 50 kolom per impor.")
                for row_index in range(sheet.nrows):
                    values = []
                    for cell in sheet.row(row_index):
                        values.append(xlrd.xldate_as_datetime(cell.value, workbook.datemode) if cell.ctype == xlrd.XL_CELL_DATE else cell.value)
                    raw.append(values)
            finally:
                workbook.release_resources()
        except ValidationError:
            raise
        except Exception as exc:
            raise ValidationError("Excel .xls tidak dapat dibaca. Coba unduh ulang sebagai Excel asli atau CSV; file HTML yang dinamai .xls belum didukung.") from exc
    else:
        raise ValidationError("Format yang didukung: .csv, .xlsx, atau .xls.")
    if not raw or any(len(row) > MAX_COLUMNS for row in raw):
        raise ValidationError("File kosong atau memiliki lebih dari 50 kolom.")
    header_index = next((i for i, row in enumerate(raw[:40]) if any(_header(cell) in ALIASES["date"] for cell in row)), 0)
    headers = [str(value or f"Kolom {i + 1}")[:120] for i, value in enumerate(raw[header_index])]
    rows = [[_cell(value) for value in row] for row in raw[header_index + 1:]]
    if len(rows) > MAX_ROWS or not rows:
        raise ValidationError("File harus berisi 1 sampai 5.000 baris data di bawah judul kolom.")
    return headers, rows


def parse_money(value, number_format):
    if isinstance(value, dict):
        text = value.get("number", "")
    else:
        text = str(value).strip().replace("\u00a0", "").replace(" ", "")
        text = re.sub(r"^(?:Rp\.?|IDR)", "", text, flags=re.I)
        if text in {"", "-"}:
            return Decimal("0.00")
        if text.startswith("(") and text.endswith(")"):
            text = "-" + text[1:-1]
        if number_format == "id":
            if not re.fullmatch(r"-?(?:\d+|\d{1,3}(?:\.\d{3})+)(?:,\d{1,2})?", text):
                raise ValidationError("Nominal tidak sesuai format Indonesia. Periksa pemisah ribuan dan desimal.")
            text = text.replace(".", "").replace(",", ".")
        elif number_format == "en":
            if not re.fullmatch(r"-?(?:\d+|\d{1,3}(?:,\d{3})+)(?:\.\d{1,2})?", text):
                raise ValidationError("Nominal tidak sesuai format internasional. Periksa pemisah ribuan dan desimal.")
            text = text.replace(",", "")
        else:
            raise ValidationError("Pilih format angka Indonesia atau internasional.")
    try:
        amount = Decimal(text)
        if not amount.is_finite() or abs(amount) > Decimal("9999999999999999.99") or amount != amount.quantize(Decimal("0.01")):
            raise ValueError
        return amount.quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError) as exc:
        raise ValidationError("Nominal tidak valid atau memiliki lebih dari dua angka desimal.") from exc


def parse_timestamp(value):
    if not isinstance(value, str):
        raise ValidationError("Tanggal tidak dikenali. Gunakan kolom tanggal bank.")
    text = value.strip()
    try:
        parsed = datetime.fromisoformat(text)
    except ValueError:
        parsed = None
        for format in ("%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d", "%d/%m/%Y %H:%M:%S", "%d %b %Y"):
            try:
                parsed = datetime.strptime(text, format)
                break
            except ValueError:
                pass
    if parsed and timezone.is_aware(parsed):
        parsed = timezone.localtime(parsed).replace(tzinfo=None)
    if not parsed or not 2000 <= parsed.year <= 2200 or parsed.date() > timezone.localdate():
        raise ValidationError("Tanggal harus lengkap dengan tahun dan tidak boleh di masa depan.")
    return parsed


def parse_date(value):
    return parse_timestamp(value).date()


def parse_reference(value):
    if isinstance(value, dict):
        try:
            number = Decimal(value.get("number", ""))
            if not number.is_finite() or number < 0 or number != number.to_integral_value() or number >= Decimal("1e15"):
                raise ValueError
            return str(int(number))
        except (ValueError, InvalidOperation) as exc:
            raise ValidationError("Referensi angka Excel tidak aman atau bukan bilangan bulat. Gunakan referensi teks asli bank tanpa kehilangan digit.") from exc
    return str(value).strip()


def normalize_statement(rows, mapping):
    def value(row, field):
        index = mapping.get(field, "")
        return row[int(index)] if index != "" and int(index) < len(row) else ""
    transactions, errors, seen = [], [], set()
    start = int(mapping.get("first_row", 1))
    end = int(mapping.get("last_row", len(rows)) or len(rows))
    if not 1 <= start <= end <= len(rows):
        raise ValidationError("Rentang baris harus berada dalam data file.")
    for number in range(start, end + 1):
        row = rows[number - 1]
        if not any(value != "" for value in row):
            continue
        try:
            timestamp = parse_timestamp(value(row, "date"))
            day = timestamp.date()
            currency = str(value(row, "currency")).strip().upper()
            if currency and currency not in {"IDR", "RP", "RUPIAH"}:
                raise ValidationError("Hanya mutasi dalam Rupiah yang didukung.")
            if mapping.get("amount", "") != "":
                amount = parse_money(value(row, "amount"), mapping["number_format"])
            else:
                debit = parse_money(value(row, "debit"), mapping["number_format"])
                credit = parse_money(value(row, "credit"), mapping["number_format"])
                if debit < 0 or credit < 0 or (debit and credit):
                    raise ValidationError("Satu baris harus debit atau kredit positif, bukan keduanya.")
                amount = credit - debit
            if not amount:
                raise ValidationError("Nominal nol bukan transaksi. Sesuaikan rentang baris untuk melewati judul/saldo awal/total.")
            description = str(value(row, "description")).strip()
            reference = parse_reference(value(row, "reference"))
            if not reference and mapping.get("reference_mode") == "balance_fingerprint":
                if mapping.get("balance", "") == "" or value(row, "balance") == "" or not description:
                    raise ValidationError("Referensi otomatis memerlukan keterangan dan saldo berjalan dari bank.")
                balance = parse_money(value(row, "balance"), mapping["number_format"])
                canonical = f"{timestamp.isoformat(timespec='microseconds')}|{description}|{amount}|{balance}"
                reference = "BNI-STMT-" + hashlib.sha256(canonical.encode()).hexdigest()[:48]
            if not reference or len(reference) > 160 or len(description) > 500:
                raise ValidationError("Referensi wajib (maks. 160 karakter), keterangan maks. 500 karakter.")
            if reference in seen:
                raise ValidationError("Referensi kembar di file. Gunakan referensi unik dari bank atau periksa transaksi yang tampak sama.")
            seen.add(reference)
            transactions.append({"date": day.isoformat(), "reference": reference, "description": description, "amount": str(amount), "row": number})
        except (ValidationError, ValueError, IndexError, TypeError) as exc:
            errors.append({"row": number, "message": " ".join(exc.messages) if isinstance(exc, ValidationError) else "Pemetaan kolom tidak valid."})
    if not transactions and not errors:
        errors.append({"row": start, "message": "Tidak ada transaksi dalam rentang ini."})
    return transactions, errors


def canonical_csv(rows):
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["date", "reference", "description", "amount"])
    for row in rows:
        writer.writerow([row[name] for name in ("date", "reference", "description", "amount")])
    return output.getvalue()
